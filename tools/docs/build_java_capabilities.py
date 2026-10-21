#!/usr/bin/env python3
"""使用工作区文档运行时生成桌面交付，不作为业务服务依赖。"""
import argparse
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from lingxi_flight_control.java_tools import TOOLS, JavaActionRequest, encode


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--output',required=True)
    parser.add_argument('--python-sha',required=True)
    parser.add_argument('--tests',required=True)
    parser.add_argument('--deployed',action='store_true')
    args=parser.parse_args()
    output=Path(args.output)
    if output.exists():
        raise SystemExit('refusing to overwrite existing document')
    doc=Document()
    section=doc.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(.7)
    section.left_margin=section.right_margin=Inches(.75)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3'):
        style=doc.styles[name]
        style.font.name='PingFang SC';style.font.color.rgb=RGBColor(0,0,0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'PingFang SC')
    doc.styles['Normal'].font.size=Pt(10.5)
    doc.styles['Normal'].paragraph_format.line_spacing=1.12
    doc.styles['Normal'].paragraph_format.space_after=Pt(5)
    doc.styles['Title'].font.size=Pt(25)
    doc.styles['Heading 1'].font.size=Pt(17)
    doc.styles['Heading 2'].font.size=Pt(12)
    for name in ('Heading 1','Heading 2'):
        doc.styles[name].paragraph_format.space_before=Pt(8)
        doc.styles[name].paragraph_format.space_after=Pt(5)
    doc.core_properties.title='大疆飞控能力与Agent工具输入输出说明'
    doc.core_properties.subject='原Java控制语义与自然语言工具契约'
    doc.core_properties.author=''

    def p(text,style=None):
        return doc.add_paragraph(text,style)
    def h(text): doc.add_heading(text,level=1)
    def page(text): doc.add_page_break();h(text)
    def code(text):
        paragraph=p(text)
        paragraph.paragraph_format.space_after=Pt(4)
        for run in paragraph.runs:
            run.font.name='Consolas';run.font.size=Pt(9)
            run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'PingFang SC')
    def table(headers,rows,widths):
        t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
        for column,width in zip(t.columns,widths):column.width=Inches(width)
        for row_values in [headers]+rows:
            row=t.rows[0] if row_values is headers else t.add_row()
            for cell,text,width in zip(row.cells,row_values,widths):
                cell.width=Inches(width);cell.text=str(text);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcPr=cell._tc.get_or_add_tcPr()
                borders=OxmlElement('w:tcBorders')
                for edge in ('top','left','bottom','right'):
                    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
                tcPr.append(borders)
                margins=OxmlElement('w:tcMar')
                for edge in ('top','left','bottom','right'):
                    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');margins.append(el)
                tcPr.append(margins)
                for para in cell.paragraphs:
                    para.paragraph_format.space_after=Pt(3)
                    for run in para.runs:run.font.size=Pt(9.5)
                if row_values is headers:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6F1');tcPr.append(shade)
                    for run in cell.paragraphs[0].runs:run.bold=True
            pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        p('')

    p('大疆飞控能力与Agent工具输入输出说明','Title')
    p('原 ktd_backend 接口核验与自然语言控制验收','Subtitle')
    p('2026年10月21日  面向天健湖草莓机库的测试准备')
    p('本说明把自然语言输入、模型工具参数、Java请求和实际触发动作逐项对齐，供开发、测试和飞手共同核对。重点是精确数值不丢失、接口语义不混用，以及“已受理”和“已执行完成”不混为一谈。')
    p('核心结论：前飞10米可以通过Java的定距分支换算成目标点；云台下调10度不能直接映射成旧云台速度接口。模型能理解角度，并不意味着现有执行接口能兑现该角度。')
    table(['证据层','本次结论'],[
        ['Java实现','核验本地origin/prod 0ed3153；候选Java中额外的跟踪改动不作为线上既有能力。'],
        ['模型工具','保留原59项通用工具契约；平台模式新增24项Java具名工具，合计83项可理解工具定义。'],
        ['执行适配','54个动作名有适配路径，包含只读状态；具体动作仍受权限、位姿、载荷和参数限制，不是54项均已实飞。'],
        ['软件验证',args.tests+'项离线测试；新增工具6条真实模型话术经隔离HTTP与合成Java验证。'],
        ['部署边界',('本文新增24项工具已更新至10086，仍须每次确认，不代用户发真机命令。' if args.deployed else
                  '10086已部署的基线为ca10023。本文新增24项工具是本次代码交付，未替换当前实例。')],
    ],[1.1,5.9])
    p('工作顺序：用户输入 → 模型提出具名工具与参数 → Agent核验资源和数值 → 用户确认 → 固定Java接口 → 记录受理或错误 → 独立核对遥测或媒体。模型不能直接调用任意URL，也不能自报另一架无人机。')
    p('全文的Java请求省略公共绑定字段。terminalId、siteId、taskId由服务端核验后补齐；令牌不交给模型。示例坐标为合成数据，不是建议飞行目标。')

    page('两个精确控制案例')
    doc.add_heading('前飞10米',level=2)
    p('输入：“往前飞10米”。预期模型输出保留方向和距离，不转换成“飞几秒”，也不改成默认步长。')
    code('move_forward({"distance_m": 10})')
    code('POST /screen/uav/operation/control\n{"type": 202, "mode": 2, "param": 10}')
    p('Java读取当前经纬度、海拔和机头航向，计算前方10米的目标坐标，再调用fly_to_point；该分支使用最大速度5米/秒。前方是机头方向，不是地图正北，也不是相机朝向。')
    p('Agent仅在同设备执行中任务、当前控制权、新鲜完整位姿及空中状态等条件满足时生成确认卡。当前相对高度低于20米等情形会被拦截，避免隐含爬升。')
    p('Java正常返回R.ok时，Agent输出submitted和device_effect=not_verified。是否真的前飞10米，需要按同一时间线核对起止坐标、航向、位置精度和设备完成回执；允许的位移误差尚未实测标定。')
    doc.add_heading('云台相对下调10度',level=2)
    p('输入：“云台向下调10度”。预期模型输出gimbal_down，angle_deg=10。“云台调低一些”需要追问角度，“调到负10度”是绝对角度，不能替换成相对下调10度。')
    code('gimbal_down({"angle_deg": 10})\n当前结果：rejected / GIMBAL_EXACT_ANGLE_UNAVAILABLE\nJava控制请求数量：0')
    p('原因：旧Java相对俯仰分支按3×stepPercentage÷zoomRatio计算速度，并调用camera_screen_drag，单位是度/秒。即使把stepPercentage设为10/3，也只是产生约负10度/秒，不是“转10度”。')
    p('param=0只停止Java的重复发送标记，未在该分支发送厂商零速度。未经角度回读、可靠停止和误差验证，不能用转动时长冒充精确角位移。')
    doc.add_heading('绝对角度也不能凭DTO注释放行',level=2)
    p('如果绕过Agent发送type=401、controlMode=1、param=-10，DTO虽然接受绝对角度字段，但该大疆实现没有通用绝对角度分支，可能直接返回R.ok而未发对应动作。Agent不会采用这种“成功但没做”的映射。')

    page('已接通的移动和航线工具')
    p('以下路径由Python执行器生成，公共绑定字段由服务端补齐。每次动作都需单独确认；旧接口没有完成回执时，不自动推进“移动完成后拍照”。')
    table(['输入与工具','Java请求核心字段','Java实际动作'],[
        ['前进10米\nmove_forward(distance_m=10)','operation/control\n202, mode=2, param=10','按机头航向计算前方目标点并fly_to_point'],
        ['后退或左右平移10米','203 / 204 / 205\nmode=2, param=10','以航向加减180或90度计算目标点'],
        ['上升或下降5米','201 / -201\nmode=2, param=5','增加或减少当前目标海拔；仍有高度门禁'],
        ['向右转15度\nturn_right(angle_deg=15)','Agent用mode=3杆量脉冲\n不是Java mode=2转向','按航向回读有界控制；目标容差和设备效果需另验'],
        ['暂停当前航线\nroute_pause()','route/control type=110','暂停航线，申请厂商飞行/载荷控制，进入DRC'],
        ['恢复航线\nroute_resume()','route/control type=112','恢复航线并退出DRC'],
        ['紧急停止\nemergency_stop()','route/control type=113','申请厂商航权，进入DRC，等待后发急停'],
        ['返航\nreturn_home()','route/control type=114\ninterrupt=false','设置任务中断标志，发返航并退出DRC'],
        ['取消返航\ncancel_return_home()','route/control type=-114','取消返航，申请航权并进入DRC'],
    ],[2.05,2.05,2.9])
    p('注意：Java的mode=1/3参数是速度或杆量百分比，不是米数。其mode=2的左右转向分支按时间循环发送，不应当作精确角度API。Agent的转向实现有独立航向回读，但没有厂商完成回执仍不能当作全面实飞验收。')

    page('已接通的相机和云台工具')
    table(['输入与工具','Java请求','实际动作与限制'],[
        ['拍张照\ncamera_photo()','camera/control type=100','申请航权和载荷权，切拍照模式，再发单拍'],
        ['开始录像或停止录像','type=101 / 102','开始前切录像模式；停止发送recording_stop'],
        ['切变焦可见光或红外或广角','type=105\ncameraMode=6 / 2 / 5','切直播镜头；可见光采用已实现的zoom镜头，不发无分支的mode=1'],
        ['光学变焦5倍\ncamera_zoom(value=5)','type=103\ncameraMode=6, param=5','设置变焦；Java DTO为Integer，非整数倍明确拒绝'],
        ['拍全景\ncamera_panorama()','type=125','先把云台朝北，再切全景模式并拍照；包含等待及平台计数副作用'],
        ['红外下限10度上限40度','type=127\njsonParams含上下限','开启等温线；整数温度、上下限角色不颠倒'],
        ['云台复位\ngimbal_reset()','gimbal/control type=403','gimbal_reset模式0'],
        ['云台垂直向下\ngimbal_90()','type=404, param=-90','gimbal_reset模式3；实际角度需回读'],
        ['云台45度俯视\ngimbal_45()','type=404, param=-45','复位后发送速度控制，草莓分支另发修正速度；不是通用绝对角度保证'],
        ['云台向下调10度','不下发旧速度接口','识别angle_deg=10后明确拒绝，不偷换成默认步进'],
    ],[2.05,1.95,3.0])
    p('相机与云台动作的Java实现常先申请厂商控制权。这是原接口内部行为；Agent仍须先核验当前平台控制权并要求用户确认，不把取权成功解释为持续航权证据。')

    items=list(TOOLS.items())
    for start in range(0,len(items),4):
        page('新增具名Java工具 '+str(start//4+1))
        p('以下工具已加入平台模型工具注册和逐次确认路径。Java实现存在不等于该架飞机及其固件已支持；结果仍按受理、失败或未知记录。')
        for name,spec in items[start:start+4]:
            doc.add_heading(spec.title,level=2)
            p('用户输入：'+spec.example)
            code('模型输出：'+name+'('+json.dumps(spec.arguments,ensure_ascii=False)+')')
            endpoint,body=encode(JavaActionRequest('example','example',name,spec.arguments,1))
            code('Java请求：'+endpoint+'\n'+json.dumps(body,ensure_ascii=False))
            p('Java动作：'+spec.effect+'。')

    page('Java已有但仍需专门适配的飞行能力')
    p('下列能力在Java中有实现，但不能直接套用现有工具的近似参数。它们记录在能力清单中，当前不作为可随意下发的模型工具；有缺失资源或高度基准时，预期输出是澄清或不可用。')
    table(['用户需求','原Java入口与字段','执行语义及缺口'],[
        ['起飞到一个位置','route/control type=117\ntakeoffPointLng/Lat/Alt','调用takeoff_to_point；高度为WGS84绝对高，不是单纯“离地30米”。工具需完整坐标、高度与起飞上下文。'],
        ['飞到指定点','fly/control type=119\nlongitude,latitude,height,speed','fly_to_point；speed在Java转整数。需要明确高度基准和速度，不能从当前简化工具补猜。'],
        ['围绕目标盘旋','fly/control type=121','DJI分支进入POI并设速度，不直接使用radius参数；不可把“半径20米”原样忽略。'],
        ['到点环绕','fly/control type=124 / -124','控制器对DJ3生成并下发环绕航线；涉及航线准备和平台持久化，不是一个简单杆量工具。'],
        ['去另一个机库降落','route/control type=123\nlandHangarCode','发指定机库返航；需核验目标机库资源和可达性，当前工具缺少目标机库绑定。'],
        ['上传航线或原地降落','route/control type=115/116/118','该大疆实现返回“不支持”。不能因为其他品牌支持，就映射到这里。'],
        ['使用速度和杆量','operation/control mode=1/3','按百分比换算并周期发送，需心跳和停止机制；与精确距离不同，不向模型开放无界持续控制。'],
    ],[1.55,2.25,3.2])

    page('喊话和机库维护能力边界')
    p('这些能力来自同一Java实现，但属于载荷播放或机库维护，不应混入本次飞行中的精确数值控制。当前自然语言飞控工具不开放下面的维护写操作。尤其格式化和关机不能由模型在普通飞行指令中顺带调用。')
    table(['喊话操作码','Java实际动作','工具接入前还需什么'],[
        ['101 / 103','重播 / 停止；T40S使用PSDK控件映射','真实载荷型号、设备位号、当前播放资源'],
        ['106','播放文件，要求文件名、URL及MD5','可信文件资源选择；不能让模型构造任意URL'],
        ['107 / 108','循环 / 单次播放','明确播放模式，当前载荷支持'],
        ['109','音量0到100','校验整数音量和目标喊话器'],
        ['201','TTS文本播报；T40S为输入框文本','用户明确的播报内容和载荷绑定'],
        ['102 / 104 / 105','返回不支持','不建立虚假的执行工具'],
    ],[1.15,2.8,3.05])
    table(['机库操作码','Java实际动作'],[
        ['100','设备重启'],['103 / 104','无人机开机 / 关机'],['105 / 106 / 113','开舱 / 关舱 / 强制关舱'],
        ['107 / 108','进入 / 退出调试模式'],['109 / 110','推杆打开 / 关闭'],['111 / 112','充电开启 / 关闭'],
        ['114 / 115','补光开关 / 空调模式'],['116 / 117','格式化机库存储 / 格式化无人机存储'],
        ['118 / 119','图传模式 / RTK坐标标定'],['101 / 102','该实现返回不支持'],
    ],[1.45,5.55])

    page('容易产生假成功的接口组合')
    table(['看起来能调用','Java实际行为','Agent预期处理'],[
        ['云台type401或402，controlMode=1，任意角度','DTO有绝对模式注释，DJI实现没有通用绝对角度分支','拒绝精确角度映射，不把R.ok当完成'],
        ['云台type404，param既不是0也不是-45/-90','返回R.ok并提示暂不支持','不注册任意绝对角度工具'],
        ['云台FPV type421','返回R.ok，提示暂不支持','不提供可执行FPV工具'],
        ['停止变焦 type104或部分未处理操作码','没有对应分支，可能落到R.ok','不凭DTO操作码清单宣称可执行'],
        ['cameraMode=1切可见光','大疆镜头切换只实现2、5、6','使用明确的6变焦或5广角，不发1'],
        ['type129被某些注释标为激光测距','本次源码常量与实现是CAMERA_LOOK_AT','按实现作为地理目标瞄准，不作为激光开关'],
        ['HTTP200或Java code200','可能只是提交，也可能没有匹配分支','分别记录submitted、错误或未知；设备完成需另证'],
    ],[2.2,2.5,2.3])
    p('云台45度是另一个需要实测的例子：名称是45度，但Java使用速度和等待近似实现。文档只能说明代码会发送什么，不能凭名称保证最终角度恰好为负45度。')

    page('模型调用与实飞验收方法')
    p('模型只选择已注册的具名工具并填写其schema。模型不直接填写terminalId、siteId、taskId、HTTP路径或原始jsonParams；这些由Agent固定映射。图像x、y、width、height须与用户明确给出的对应值一致，矩形不能超出0到1范围。')
    p('新增工具沿用已有会话、确认绑定、SQLite幂等账本和审计。任务结束、控制权变化、设备绑定不符、载荷缺失或确认过期时拒绝；不自动重发未知结果，不在后台自动发真机命令。原模拟器工具面保持59项，避免把平台专用能力伪装成模拟成功。')
    table(['观察阶段','预期证据'],[
        ['理解','用户话术的摘要、工具名、明确参数；含糊输入先澄清'],
        ['确认','卡片说明动作、设备与参数；未确认时Java写请求数为0'],
        ['提交','同一command_id/trace_id关联Java请求路径、操作码和数值'],
        ['接口返回','记录HTTP状态和Java业务码，不输出凭据或原始响应正文'],
        ['设备完成','实际位移、航向、云台角度或关联照片；缺少这些则标记缺证据'],
    ],[1.2,5.8])
    p('下午建议逐条测试：先单独拍照，核对媒体；再在合适高度和任务条件下测试前飞10米，核对位移；新增相机设置逐项确认并检查实际画面。云台精确下调10度仍以明确拒绝为预期，不进行“看起来差不多”的验收。')
    p('审计工具tools/flight_audit.py按任务ID和时间窗口只读查询。旧定时监控保持暂停；用户开始实飞时再绑定本轮任务。当前审计没有自动采集厂商遥测和媒体，因此日志完整不等于设备效果已被完整验证。')

    page('版本和来源索引')
    p('Java源码核验基线：本地origin/prod 0ed3153fd0f4fafd6fa1f0513cb5deb1d20b9826。Java仓库本轮没有修改。线上Java是否完全等于该提交，尚未通过部署清单核验。')
    p('Python工具实现提交：'+args.python_sha+'。'+('10086已部署此版本。' if args.deployed else
        '10086现行运行版本ca100239f609db789cf33f6783d67e0e64479ac3；本次新增工具不与该已部署版本混记。'))
    table(['源码文件','核验内容'],[
        ['LargeScreenController.java','固定接口、任务/设备路由、参数校验及平台副作用'],
        ['DJHangarOperation.java','航线、位移、云台、相机、喊话及机库实际分支'],
        ['DJUtil.java','厂商方法名称、速度单位、椭球高与消息类型'],
        ['UavOpeConstant.java\nLoadOpeConstant.java','操作码；以实现分支复核，不单独作为支持证明'],
        ['CameraControlDTO.java\nGimbalControlDTO.java\nTakeoffToPointDTO.java','字段类型、参数取值与注释；识别注释和实现不一致'],
        ['java_tools.py\nktd_platform.py\nbrowser_gateway.py','本次24项具名工具、固定编码、授权绑定与确认'],
        ['test_java_tools.py\ndji_browser_contract.py','参数边界、逐次确认、幂等回放与真实模型隔离HTTP验证'],
    ],[2.4,4.6])
    p('验证结论：'+args.tests+'项离线回归通过；前端150项测试与构建通过；新增6条真实模型输入均生成预期工具并在确认后提交到合成Java，确认前0请求、真机请求0。未在本轮证明任何新能力的实际设备效果。')
    output.parent.mkdir(parents=True,exist_ok=True)
    doc.save(output)
    print(output)


if __name__=='__main__':main()
